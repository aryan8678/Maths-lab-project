"""Build the mini thesis and attach every R drawing in a separate atlas."""
from pathlib import Path
import json,csv,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'submission';OUT.mkdir(exist_ok=True)
D=json.loads((ROOT/'.build/content.json').read_text(encoding='utf-8'))
def rows(p):
 with (ROOT/p).open(newline='',encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def f(x,d=2):return f'{float(x):.{d}f}'
def p(x):return '<0.001' if float(x)<.001 else f'{float(x):.3f}'
def percent(x):return f'{100*float(x):.2f}%'
M=D['metrics'];A=M['setup'];O=M['outliers'];C=M['correlation'];R=M['rbd'];L=M['latin'];N=M['fits'];B=M['binomial'];V=M['anova'];P=M['paired']
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.68);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Caption'].font.size=Pt(9)
header=sec.header.paragraphs[0];header.text='BMAT202P Probability and Statistics Lab';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Page ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title=D['title'];doc.core_properties.author='; '.join(D['authors'])
def para(t):return doc.add_paragraph(t)
def page(title):doc.add_page_break();doc.add_heading(title,level=1)
def sub(title):doc.add_heading(title,level=2)
def table(values,widths=None):
 t=doc.add_table(rows=0,cols=len(values[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(values):
  cells=t.add_row().cells
  for j,val in enumerate(row):
   c=cells[j];c.text=str(val);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if widths:c.width=Inches(widths[j])
   props=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   props.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173D54' if i==0 else ('F0F4F6' if i%2==0 else 'FFFFFF'));props.append(shade)
   margin=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    v=OxmlElement('w:'+edge);v.set(qn('w:w'),'80');v.set(qn('w:type'),'dxa');margin.append(v)
   props.append(margin)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(3);par.paragraph_format.space_before=Pt(3)
    par.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in par.runs:
     run.font.size=Pt(10);run.bold=i==0;run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
  trp=t.rows[-1]._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
  if i==0:trp.append(OxmlElement('w:tblHeader'))
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
figure_no=0
def figure(path,caption,height=2.7):
 global figure_no
 figure_no+=1
 par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.keep_with_next=True
 im=Image.open(ROOT/path);w,h=im.size;width=min(6.9,height*w/h)
 par.add_run().add_picture(str(ROOT/path),width=Inches(width))
 cap=doc.add_paragraph(f'Figure {figure_no}  {caption}',style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
def fp(folder,name):return folder+'/output/'+name+'.png'
def equation(t):
 par=doc.add_paragraph();om=OxmlElement('m:oMath');run=OxmlElement('m:r');v=OxmlElement('m:t');v.text=t;run.append(v);om.append(run);par._p.append(om)
doc.add_paragraph('Statistical analysis of diabetes and lifestyle indicators',style='Title')
doc.add_paragraph('Mini project report',style='Subtitle')
para('BMAT202P Probability and Statistics Lab\nR programming and the prescribed statistical experiments')
for name in D['authors']:para(name)
sub('Abstract')
para('We analyse a reproducible random sample of 3,000 respondents from the processed CDC BRFSS 2015 diabetes health-indicators dataset. The project describes BMI and health status, examines relationships with lifestyle variables, fits probability distributions and applies the hypothesis tests prescribed in the course. Published experiments supplement the survey for paired measurements and randomized designs.')
para('The sample has 469 respondents with prediabetes or diabetes. BMI and poor physical health days have a weak positive correlation (r = 0.102). The simple regression explains 1.03% of BMI variation and the multiple model explains 6.81%. BMI filtering flags 114 records and changes both descriptive estimates and the positive-status proportion. The normal and Poisson models fit poorly. Experimental supplements illustrate genuine pairing, completely randomized assignment, complete blocks and a Latin Square.')
para('The findings are exploratory associations. Self-report, historical data, omitted survey weights, confounding and uneven group sizes limit generalization. This report records methods, assumptions and uncertainty alongside the results.')
page('Aim and dataset')
para('Our aim is to describe BMI and health indicators, examine their associations and demonstrate every statistical method in the BMAT202P syllabus. We ask how BMI varies across health groups, how much variation the regression models explain and whether the observations follow the proposed probability distributions. We use separate experimental datasets where the cross-sectional survey cannot support a required design.')
sub('Data source and collection')
para('The original CDC BRFSS 2015 survey contains 441,456 records and uses survey weighting. We use Alex Teboul’s processed Kaggle derivative with 253,680 rows and 21 predictors plus a three-class target, giving 22 raw columns. This is secondary data collection. Acquisition verifies the exact source file with a SHA-256 checksum saved in data/provenance.json. [1–3]')
table([['Dataset','No diabetes','Prediabetes','Diabetes'],['Full derivative']+[r['count'] for r in rows('Lab00_Setup/output/class_balance.csv')[:3]],['Working sample','2,531','48','421']],[2.3,1.55,1.55,1.6])
para(f"We select 3,000 rows without replacement using set.seed(101). Combined-positive status means prediabetes or diabetes. Its proportion is {percent(A['sample_combined_rate'])} in the working sample and {percent(A['full_combined_rate'])} in the derivative. Similar proportions are a useful sampling check, but do not prove representativeness. These unweighted proportions are not national prevalence estimates.")
sub('Experimental supplements')
table([['Concept','Published dataset','Records'],['Paired t test','MASS anorexia, CBT group','29 pairs'],['CRD','Cochran potato sulfur experiment','32 plots'],['RBD','Gomez rice seeding rates','24 plots'],['Latin Square','Goulden wheat fungicide experiment','25 plots']],[1.25,4.4,.95])
para('The agricultural tables are extracted from agridat version 1.26 with their documented units and design. They demonstrate statistical designs and are not pooled with the health data. [4–7]')
page('Variable definitions and data preparation')
table([['Variable','Definition and handling'],['Diabetes_012','Three categories: no diabetes, prediabetes, diabetes'],['Diabetes_binary','Yes for prediabetes OR diabetes; No otherwise'],['BMI','Body mass index, kg/m², recorded as integer values'],['MentHlth and PhysHlth','Self-reported poor-health days, integer 0–30'],['Age','13 age bands; categorical contrasts in regression'],['Income and Education','Ordinal bands; categorical contrasts in regression'],['Smoker','Ever smoked at least 100 cigarettes'],['PhysActivity','Physical activity in past 30 days, excluding job']],[1.75,5.25])
para('Preparation checks required columns, allowed categorical codes, missing values, finite positive BMI and day counts in the range 0–30. No invalid or missing records occurred in this working sample. Records retain their original source-row identifiers. The prepared RDS preserves factor levels, while CSV provides a portable exchange format. The complete dictionary is in Lab01_Data_Types/output/variable_dictionary.csv.')
para('Thirteen sample records repeat an observed feature profile. We retain these profiles because separate respondents can have identical answers. The predictors do not provide a unique person identifier. Removing identical feature rows would assume duplicate people without evidence.')
para('The master runner executes the individual lab scripts in dependency order. It creates output directories and saves tables, figures and session information. The analysis uses base R and the recommended MASS package, keeping the statistical workflow close to the course.')
sub('Traceable analysis datasets')
table([['Stage','Rows','Decision'],['Working sample','3,000','Seeded sampling without replacement'],['Validated primary sample','3,000','Keep valid extreme BMI values'],['BMI filtered sensitivity','2,886','Remove 114 BMI IQR flags']],[2.2,.7,4.1])
page('Outlier analysis and removal')
para('We calculate BMI quartiles and use the conventional 1.5 IQR fences. Q1 = 24, Q3 = 31 and IQR = 7, giving lower and upper fences of 13.5 and 41.5. The sample ranges from BMI 15 to 89. The rule flags 114 records, all above the upper fence, or 3.8% of the sample. The ledger records each source row, BMI and diagnostic category.')
figure(fp('Lab01B_Outlier_Analysis','bmi_outlier_fences'),'BMI distribution with the calculated outlier fences',2.3)
para('We retain valid extremes in the primary analysis. To demonstrate removal, we create a separate dataset that excludes all BMI flags. This restriction changes the population being analysed. Binary and ordinal codes use validation rules instead of IQR filtering. Zero-heavy day variables can have small or zero IQR and should not be automatically filtered.')
table([['Measure','Primary sample','Filtered sensitivity'],['n','3,000','2,886'],['Mean BMI','28.40','27.65'],['BMI SD','6.35','4.96'],['BMI variance','40.34','24.59'],['Combined-positive proportion','15.63%','14.62%'],['BMI and PhysHlth correlation','0.102','0.056']],[2.8,2.1,2.1])
para('Removal lowers mean BMI and the combined-positive proportion. These changes show why a valid extreme value should not automatically be treated as an error. The original dataset remains available and each comparison states its analysis sample.')
page('Descriptive and graphical analysis')
table([['Variable','Mean','Median','SD','Variance']]+[[r['variable'],f(r['mean']),f(r['median'],0),f(r['sd']),f(r['variance'])] for r in D['tables']['summary']],[2,1.25,1.25,1.25,1.25])
para('BMI has mean 28.40 and median 27, consistent with a right tail. Mean poor mental-health days is 3.26 and mean poor physical-health days is 4.15. Both count variables have median zero and large standard deviations. Variance uses squared units: (kg/m²)² for BMI and days² for the health-day counts.')
figure(fp('Lab02_Summary_Stats','box_bmi_by_diabetes'),'BMI by the original three diagnostic categories',2.8)
table([['Health status','n','Mean BMI','95% CI'],['No diabetes','2,531','27.72','27.49 to 27.95'],['Prediabetes','48','32.85','30.71 to 35.00'],['Diabetes','421','31.98','31.29 to 32.68']],[2.1,.8,1.2,2.9])
para('Mean BMI is higher in the prediabetes and diabetes groups, but the individual distributions overlap. The prediabetes estimate has greater uncertainty because the group is small. Lab 2 also produces a histogram, education bar chart, sex pie chart, BMI density and a discrete day-count plot. The complete figure atlas retains every drawing.')
page('Correlation and simple linear regression')
para('Pearson correlation measures linear association between BMI and PhysHlth. We obtain r = 0.1016, with a 95% confidence interval from 0.0661 to 0.1369 and p < 0.001. Spearman’s rank correlation is 0.1079. Both indicate a weak positive association in the working sample.')
matrix=rows('Lab03_Correlation_Regression/output/correlation_matrix.csv')
table([['Pearson r','BMI','MentHlth','PhysHlth']]+[[r['variable']]+[f(r[v],3) for v in ['BMI','MentHlth','PhysHlth']] for r in matrix],[2.2,1.6,1.6,1.6])
para('Covariance measures shared variation in the original units. Its diagonal equals each variable’s variance. Because it depends on scale, correlation is easier to compare across variables. The covariance matrix is saved separately with the same analysis rows.')
figure(fp('Lab03_Correlation_Regression','scatter_bmi_physhlth'),'BMI and poor physical-health days with the fitted regression line',2.7)
equation('Predicted BMI = 28.0874 + 0.0749 × PhysHlth')
para('The slope means that one additional poor physical-health day is associated with about 0.075 BMI units on average in the fitted line. Its 95% CI is 0.0486 to 0.1012. R² = 0.0103, so this predictor accounts for only 1.03% of observed BMI variation. Statistical significance does not make this a strong predictor or a causal relationship.')
page('Multiple regression and model assessment')
para('The multiple linear model uses BMI as the response and Age, PhysActivity, Income, Education and Smoker as predictors. Age, income and education use categorical contrasts, avoiding an assumption of equally spaced numerical bands. Reference categories are age 50–54, income 35–50k, high school, inactivity and never having smoked 100 cigarettes.')
table([['Predictor','df','Partial F','p']]+[[r['term'],r['Df'],f(r['F.value']),p(r['Pr..F.'])] for r in D['tables']['joint']],[2.8,.7,1.5,2])
para('The model has R² = 6.81% and adjusted R² = 5.99%. Age, physical activity, income and education have statistically detectable joint associations with BMI after adjustment. The smoking coefficient is −0.063 BMI units, with CI −0.521 to 0.394 and p = 0.786. This BMI question differs from the unadjusted status-by-smoking comparison.')
figure(fp('Lab04_Multiple_Regression','multiple_diagnostics'),'Residual shape, spread and leverage diagnostics for the multiple model',2.9)
para('Residuals and influence plots show the limits of the linear approximation. Cook’s distance greater than 4/n flags 140 observations for review, not automatic deletion. The filtered model has adjusted R² = 4.89%. Full coefficients, confidence intervals, influence rankings and coding-dependent design-column VIFs are saved as CSV. The models describe BMI associations and do not measure diabetes classification accuracy.')
page('Normal distribution fitting')
para('We compare BMI with a normal distribution fitted at mean 28.40 and SD 6.35. The overlay and Q-Q plot show a long upper tail. The Shapiro-Wilk test on all 3,000 BMI observations gives W = 0.8969 and p < 0.001, rejecting the normal teaching hypothesis.')
figure(fp('Lab06_Normal_Poisson_Fitting','normal_bmi_hist'),'BMI histogram with the fitted normal density',3)
figure(fp('Lab06_Normal_Poisson_Fitting','normal_bmi_qq'),'Normal Q-Q plot showing the upper-tail departure',2.4)
para('A normality test on raw BMI does not by itself decide whether every t-test or regression is usable. Each procedure has its own assumptions about groups, means or residuals. We retain the course procedures, check their assumptions and use appropriate sensitivity analyses where needed.')
page('Binomial and Poisson distribution fitting')
sub('Binomial model')
para('We form 300 artificial blocks of ten respondents and count combined-positive statuses. The ungrouped proportion is 0.1563. Sparse expected tail categories require pooling. We fit the parameter to the pooled frequencies and use df = number of bins − 1 − number of fitted parameters. The five-bin test gives χ² = 0.287, df = 3 and p = 0.962. It does not reject the binomial teaching model. This does not prove that all respondents share the same health risk. [8]')
figure(fp('Lab05_Binomial_Fitting','binomial_observed_vs_expected'),'Observed and fitted binomial block frequencies',2.2)
sub('Poisson model')
para('PhysHlth has mean 4.15 and variance 74.24, giving a dispersion ratio of 17.88, compared with one for Poisson. Zero days accounts for 62.9% and 30 days for 7.4% of the sample. The ordinary Poisson support is unbounded, unlike this 30-day measure. We include the theoretical upper tail and pool sparse categories. The grouped fit gives χ² = 41,781.7, df = 6 and p < 0.001. The model is inadequate.')
para('The ordinary overlay uses the raw-mean estimate λ = 4.152. The formal pooled-bin test refits λ by grouped maximum likelihood, giving 1.797. These are explicitly different estimation procedures. The minimum fitted expected count is 7.65, and the discrepancy remains large. The fit demonstrates model criticism rather than forcing agreement with the data.')
page('One sample and two sample hypothesis tests')
para('For the course demonstrations we state two-sided null and alternative hypotheses and use α = 0.05. The BMI benchmark 27 and combined-positive proportion benchmark 11% are hypothetical classroom values. They are not external national estimates. We present confidence intervals and effect magnitudes alongside p-values.')
table([['Question or comparison','Estimate and 95% CI','p'],['Mean BMI versus 27','28.40 [28.17, 28.63]','<0.001'],['Combined-positive proportion versus 11%','15.63% [14.36%, 16.99%]','<0.001'],['Female minus male BMI','−0.65 [−1.10, −0.20]','0.004'],['Never minus ever-smoking positive proportion','−3.27 pp [−5.95, −0.59]','0.016'],['Female minus male PhysHlth','+1.25 days [0.63, 1.86]','<0.001']],[3.1,2.9,1])
para('The mean comparisons use Welch’s t-test, allowing unequal group variances. The proportion tests use the explicit positive-event counts and continuity correction. The reported statistic for prop.test is Pearson X-squared, not z. The independent PhysHlth test fulfils the independent-samples part of Experiment 9.')
figure(fp('Lab08_TwoSample_Tests','bar_diabetes_by_smoker'),'Combined-positive proportion by ever-smoking status',2.8)
para('Never-smoking respondents have a positive proportion of 14.09% (223/1,583) and ever-smoking respondents 17.36% (246/1,417). These unadjusted differences can reflect confounding. Bounded and skewed day counts require caution when interpreting their mean comparison. Repeated exploratory testing increases the chance of finding small p-values.')
page('Paired t test supplement')
para('The BRFSS survey has no before-and-after observations to pair. We use the CBT group in MASS::anorexia, which contains 29 young female patients measured before and after the study period. The paired test analyses each patient’s Postwt minus Prewt difference. The source records weights in pounds. [4]')
figure(fp('Lab09_Paired_TTest','paired_weights'),'Each line connects the same patient before and after the study period',3)
table([['Quantity','Result'],['Mean before weight','82.69 lb'],['Mean after weight','85.70 lb'],['Mean paired change','+3.01 lb'],['95% CI for mean change','+0.23 to +5.79 lb'],['Paired t statistic and df','t = 2.216, df = 28'],['Two-sided p','0.035']],[3.4,3.6])
para(f"The difference distribution departs from normality (Shapiro p = 0.008). A signed-rank sensitivity test gives p = {p(P['wilcoxon_sensitivity_p'])}, but addresses a related distributional hypothesis rather than the identical mean hypothesis. Interpret the t-test cautiously. An uncontrolled change within one treatment group does not by itself prove treatment efficacy. Pairing improves the comparison by accounting for each individual’s baseline.")
page('Chi square goodness of fit and contingency tests')
para('The age-band goodness-of-fit exercise tests exactly equal probabilities across all 13 categories. The result is χ² = 678.12, df = 12 and p < 0.001. This is a hypothetical teaching null. Unequal age-band widths give no reason to expect equal population probabilities or to use this test as evidence of survey representativeness.')
figure(fp('Lab10_ChiSquare_Tests','stacked_diabetes_by_highbp'),'Combined health status within blood-pressure groups',3)
table([['Association','Chi square','Cramer V','p'],['Combined status and HighBP','226.28','0.275','<0.001'],['Combined status and Smoker','6.07','0.045','0.014']],[3.1,1.35,1.35,1.2])
para('All expected contingency counts exceed five. Pearson contingency tests use no continuity correction. Combined-positive status is 6.93% without high BP and 27.06% with high BP. Cramer’s V shows a stronger association for HighBP than for ever-smoking. A categorical association does not demonstrate that changing a predictor would change health status.')
para('The smoking contingency test and the two-proportion test address closely related equality questions, but their small p-value difference reflects the continuity correction used in the proportion test. The tests use the same event definition.')
page('One way ANOVA and completely randomized design')
sub('Observational education comparison')
para('The education ANOVA gives F = 14.41 and p < 0.001, with eta-squared 0.0235. Education groups account for about 2.35% of BMI variation in this unadjusted comparison. The variance check gives p < 0.001 and the no-school group has only two respondents. Welch ANOVA remains significant (p = 0.000351). Separate-variance pairwise comparisons with Holm adjustment support interpretation. Standard ANOVA and Tukey outputs remain available for the syllabus exercise.')
sub('Genuine randomized supplement')
para('The Cochran sulfur experiment uses completely randomized assignment to seven treatments on 32 potato plots. The untreated control O has eight replicates and six sulfur treatments have four each. Infection is the percentage of potato surface covered by scab. [5]')
figure(fp('Lab11_ANOVA_CRD','crd_infection'),'Infection percentages in the completely randomized sulfur experiment',2.7)
para('The treatment ANOVA gives F = 3.608 and p = 0.010. Tukey’s control O versus fall 1,200-lb/acre treatment F12 contrast is 16.875 percentage points, with CI 3.74 to 30.01 and adjusted p = 0.006. The residual normality check gives p = 0.872 and variance check p = 0.116, with diagnostics available in the atlas. These checks do not prove all assumptions, but do not show a clear departure at the chosen level.')
para('The educational BMI groups were not randomly assigned. We distinguish that observational one-way ANOVA from the genuine CRD supplement. Causal language requires the experimental design and its scope, not just an ANOVA command.')
page('Randomized block design')
para('The Gomez rice experiment has six seeding rates, each appearing once in each of four complete blocks. The 24 observations record seeding rate and yield in kg/ha. We verify the treatment-by-block table, then fit yield = block + rate + error. This dataset provides a straightforward RBD matched to the syllabus. [6]')
figure(fp('Lab12_ANOVA_RBD','rbd_yield_by_rate'),'Yield at the six seeding rates within each block',3)
table([['Source','df','F','p'],['Block','3','5.86','0.007'],['Seeding rate','5','2.17','0.113'],['Residual','15','','']],[2.8,1.1,1.4,1.7])
para('Residual mean square falls from 200,152 without blocks to 110,558 with blocks, a reduction of 44.8%. Comparing mean square as well as sum of squares accounts for the changed residual degrees of freedom. Block differences are detectable, while the overall rate effect is not significant at α = 0.05. The largest observed mean is at 75 kg/ha, but the non-significant overall result cannot establish an optimal rate.')
para('Blocking accounts for one nuisance source of variation. The additive model assumes no treatment-by-block interaction. The RBD diagnostic sheet and Tukey table are available for questions. We do not claim that a non-significant test proves every rate is equally effective.')
page('Latin Square design')
para('Goulden’s 5 by 5 wheat experiment tests five fungicide treatments. Each treatment occurs once in every row and every column. We verify both balances and fit yield = row + column + treatment + error. The model controls two nuisance factors simultaneously. [7]')
figure(fp('Lab13_Latin_Square','latin_square_layout'),'Latin Square treatment layout with observed yields',3)
table([['Source','df','p'],['Row','4','0.013'],['Column','4','0.263'],['Treatment','4','<0.001'],['Residual','12','']],[3.4,1.4,2.2])
para('The treatment F statistic is 21.03 with p = 0.0000237. Weekly dusting C has mean yield 13.12, compared with A 6.84, B 6.46, D 7.96 and untreated E 4.92. Tukey comparisons find C higher than each other treatment at the 5% familywise level. The source does not specify yield units, so we retain the source scale.')
para('A means dusting before rains, B after rains, C weekly dusting, D weekly drifting and E no dusting. The additive Latin Square model assumes no interactions between treatments and nuisance factors. The small experiment supports conclusions about this experiment’s conditions, not a universal farming recommendation.')
page('Consolidated findings limitations and conclusion')
sub('Answers to the project aim')
para('The sample shows higher mean BMI in the prediabetes and diabetes groups, with overlapping distributions. Poor physical-health days have a weak positive BMI association. The simple model explains 1.03% of BMI variation and the multiple model explains 6.81%, so these predictors leave most variation unexplained.')
para('Outlier removal materially changes BMI spread and health-status composition. Keeping valid extremes in the primary sample and showing a separate filtered sensitivity analysis makes the effect transparent. A boxplot flag alone is not evidence of erroneous data.')
para('The binomial block model is not rejected after a corrected pooled goodness-of-fit calculation. The normal BMI and Poisson day-count models show clear inadequacy. An unsuitable fit is a useful statistical finding. The probability exercises require checking assumptions rather than merely producing p-values.')
para('The experimental supplements demonstrate paired differences and CRD, RBD and Latin Square analysis. The RBD example shows a useful block adjustment even though the rate effect is non-significant. The Latin experiment finds a treatment difference after accounting for rows and columns.')
sub('Limitations')
para('The health dataset is self-reported, historical and observational. The derivative omits the weights and survey-design information required for national prevalence inference. A simple random sample from the derivative does not recreate the original population survey. Confounding, class imbalance and a small prediabetes group constrain interpretation. Some education categories are very sparse. Multiple exploratory tests increase false-positive risk; post-hoc comparisons use stated adjustments, but we do not claim a confirmatory health study.')
para('Regression residuals, influence and sensitivity show approximate models with limited explanatory strength. Mean tests on skewed counts and the non-normal paired changes require caution. Hypothetical benchmarks illustrate the syllabus and are not independent evidence of a public-health discrepancy.')
sub('Conclusion')
para('The project applies the complete course sequence to real data, explains why each procedure was chosen and presents uncertainty alongside results. Its health conclusions remain descriptive. Better population inference would require the original survey design, documented reference values and a pre-specified analysis. The experimental design supplements remain separate from the health conclusions.')
page('References and reproducibility')
refs=[
 ('1','CDC','2015 BRFSS Survey Data and Documentation',D['sources']['cdc']),
 ('2','Alex Teboul','Diabetes Health Indicators Dataset',D['sources']['kaggle']),
 ('3','UCI','CDC Diabetes Health Indicators variable definitions',D['sources']['uci']),
 ('4','MASS','Anorexia data from Hand et al 1993',D['sources']['paired']),
 ('5','agridat','Cochran and Cox 1957 sulfur CRD',D['sources']['crd']),
 ('6','agridat','Gomez and Gomez 1984 seeding rate RBD',D['sources']['rbd']),
 ('7','agridat','Goulden 1952 Latin Square',D['sources']['latin']),
 ('8','NIST','Chi Square Goodness of Fit Test',D['sources']['gof'])]
for no,author,title,url in refs:para(f'[{no}] {author}. {title}.\n{url}')
para('Course source: supplied BMAT202P Probability and Statistics Lab syllabus, version 1.0. Prescribed textbook: Joseph Schmuller, Statistical Analysis with R, Wiley, 2017. Evaluation source: attached four-category mini-project rubric. Sources accessed on 4 October 2026.')
para('Execution: R 4.5.3 on Windows, base packages and MASS 7.3-65. The package-extracted supplements are pinned to agridat 1.26. Sampling seed is 101. See README.md and RUNNING_GUIDE.md for the exact acquisition and execution commands. All R source files remain separate by exercise.')
page('Appendix course coverage and supporting evidence')
table([['Syllabus experiment','Implemented evidence'],['1 Data types and import export','Labs 0 and 1, dictionary and validated RDS'],['2 Summary and graphical representations','Lab 2, six plot types and frequency tables'],['3 Correlation and simple regression','Lab 3, slope, correlation CI, covariance and R²'],['4 Multiple regression','Lab 4, coefficients, joint tests, adjusted R² and diagnostics'],['5 Binomial distribution','Lab 5, block frequencies and corrected pooled GOF'],['6 Normal and Poisson distributions','Lab 6, overlays, Q-Q and support-aware pooled GOF'],['7 One sample mean and proportion','Lab 7, labelled benchmark hypotheses and CI'],['8 Two sample means and proportion','Lab 8, Welch comparison and positive-event counts'],['9 Independent and dependent t tests','Lab 9a BRFSS, Lab 9b genuine pre/post pairs'],['10 Chi square GOF and contingency','Lab 10, uniform-null GOF, expected counts and Cramer V'],['11 CRD RBD and Latin Square ANOVA','Labs 11–13, real designs, balance checks and follow-ups']],[3.0,4.0])
para('The presentation-specific outlier requirement is implemented in Lab 1B. The report/consolidated_tests.csv table supplies a common inference record. Each lab output folder includes its tables and plots. All 29 drawings are embedded in Complete_Figure_Atlas.pdf and supplied as original PNGs in All_Drawings.zip. The atlas also includes the additional diagnostic sheets for oral questions.')
para('The four rubric areas have corresponding evidence: source and dictionary for Collection of Data; justified methods, tables and graphics for Analysis; findings, uncertainty and limitations for Review; organized slides, this report, references and reproducible code for Final Submission. This mapping demonstrates coverage and does not predict awarded marks.')
doc.save(OUT/'Diabetes_Statistics_Report.docx')

# Figure atlas: preserve every R drawing at its original aspect ratio.
figures=rows('Lab14_Consolidation/output/figure_inventory.csv')
atlas=canvas.Canvas(str(OUT/'Complete_Figure_Atlas.pdf'),pagesize=(842,595))
atlas.setTitle('Complete R figure atlas');atlas.setAuthor('; '.join(D['authors']))
for i,r in enumerate(figures,1):
 path=r['figure'];name=Path(path).stem.replace('_',' ')
 atlas.setFont('Helvetica-Bold',18);atlas.drawString(32,557,f'{i:02d}  {name}')
 atlas.setFont('Helvetica',9);atlas.drawString(32,536,path)
 im=Image.open(ROOT/path);w,h=im.size;scale=min(778/w,465/h)
 atlas.drawImage(ImageReader(im),32+(778-w*scale)/2,45+(465-h*scale)/2,width=w*scale,height=h*scale)
 atlas.setFont('Helvetica',9);atlas.drawString(32,20,'R output from the stated lab. Sample definitions and interpretation are in the mini project report.');atlas.showPage()
atlas.save()
with zipfile.ZipFile(OUT/'All_Drawings.zip','w',zipfile.ZIP_DEFLATED) as z:
 for r in figures:z.write(ROOT/r['figure'],r['figure'])
 z.write(OUT/'Complete_Figure_Atlas.pdf','Complete_Figure_Atlas.pdf')
print('Built mini thesis, 29-page complete figure atlas and original PNG drawing archive.')
